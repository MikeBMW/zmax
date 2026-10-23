#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成脚本 vs 交付文档 的一致性闸。

为什么需要：把生成脚本改了几十处之后再从别处 `cp` 覆盖它，产物是对的、脚本是旧的，
下次重跑会把所有版式修正静默回退。这个脚本把两份 pptx 的每页每个 shape 的
（名称/坐标/尺寸/文本/表格列宽）做成签名逐项比对，差异 0 才允许相信「脚本能重现产物」。

用法:
    python verify_deck_signature.py regen.pptx shipped.pptx   # 差异 0 → 退出码 0
    python verify_deck_signature.py shipped.pptx             # 只打印签名摘要

需要在有 python-pptx 的解释器下跑（找不到就换带它的 venv）。
"""
import sys

from pptx import Presentation

EMU = 914400.0


def sig(path):
    prs = Presentation(path)
    pages = []
    for slide in prs.slides:
        items = []
        for sh in slide.shapes:
            txt = sh.text_frame.text if sh.has_text_frame else ""
            if getattr(sh, "has_table", False) and sh.has_table:
                txt = "|".join(c.text for row in sh.table.rows for c in row.cells)
                txt += "|w:" + ",".join("%.2f" % (c.width / EMU) for c in sh.table.columns)
            items.append((sh.name,
                          round((sh.left or 0) / EMU, 2),
                          round((sh.top or 0) / EMU, 2),
                          round((sh.width or 0) / EMU, 2),
                          round((sh.height or 0) / EMU, 2),
                          txt))
        pages.append(items)
    return pages


def compare(a_path, b_path, verbose=8):
    a, b = sig(a_path), sig(b_path)
    print("页数: %d vs %d" % (len(a), len(b)))
    diff = 0
    for i, (sa, sb) in enumerate(zip(a, b), 1):
        if len(sa) != len(sb):
            print("P%d shape 数不同: %d vs %d" % (i, len(sa), len(sb)))
            diff += 1
            continue
        for ja, jb in zip(sa, sb):
            if ja != jb:
                diff += 1
                if diff <= verbose:
                    print("P%d 差异:" % i)
                    for x, y in zip(ja, jb):
                        if x != y:
                            print("   A=%s" % (str(x)[:100],))
                            print("   B=%s" % (str(y)[:100],))
    print("差异单元: %d" % diff)
    return diff


if __name__ == "__main__":
    if len(sys.argv) == 3:
        sys.exit(0 if compare(sys.argv[1], sys.argv[2]) == 0 else 1)
    if len(sys.argv) == 2:
        for i, page in enumerate(sig(sys.argv[1]), 1):
            print("P%-2d shapes=%d" % (i, len(page)))
        sys.exit(0)
    print(__doc__)
    sys.exit(2)
