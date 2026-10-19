#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""荣旗立项资料清单 → xlsx (17 项, 客户可直接填)。

真源口径: 与 docs/outreach/荣旗_立项资料清单_邮件_20261019.md 的附表逐行一致。
列: 序号 / 资料项 / 为什么需要(用途) / 期望格式 / 期望时间 / 贵方负责人 / 贵方回复
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "docs", "outreach", "荣旗_立项资料清单_20261019.xlsx")

ROWS = [
    ("一、场景与需求", None, None, None, None),
    (1, "工件型号/尺寸/材质/易损点、是否多料号混线", "决定夹爪与视觉方案、换型时间", "规格书或表格", "10-21"),
    (2, "来料方式（Tray/散料/管装/编带）、单盘数量、换盘方式", "决定上料机构与节拍瓶颈", "文字+照片", "10-21"),
    (3, "计划替代的工位数量、布局尺寸、上下料方式（人/AGV）", "决定单台覆盖工位数与销量预测", "布局图/草图", "10-21"),
    (4, "现有节拍（秒/颗 或 UPH）、每班人工编制、班次", "算 ROI 与节拍目标", "数字", "10-21"),
    (5, "供电/气源/占地/洁净度/ESD 等级/是否允许改造", "现场可实施性评估", "文字", "10-21"),
    (6, "MES/PLC 对接需求与需回传字段", "决定软件工作量", "接口清单", "10-22"),
    ("二、客户验收指标（写进合同判据）", None, None, None, None),
    (7, "成功率/良率目标值与统计口径（按颗/按班）", "合同验收判据", "客户原话", "10-21"),
    (8, "节拍目标值，是否接受分阶段爬坡（PVT 达标/量产提优）", "方案指标与里程碑", "客户原话", "10-21"),
    (9, "定位重复性、插入深度/横向偏差允许范围", "精度设计与标定工作量", "数字+口径", "10-21"),
    (10, "uptime 或 MTBF 要求、连续运行验收时长", "可靠性设计与验收", "数字", "10-22"),
    (11, "验收阶段（EVT/DVT/PVT）判据、抽样规则、验收地点", "立项文档验收章节", "文字", "10-22"),
    ("三、商务与计划", None, None, None, None),
    (12, "2026/2027/2028 年需求台数与扩产依据", "销量预测（立项必填）", "表格", "10-21"),
    (13, "目标预算区间或投资回收期要求", "配置选型与报价区间（我方不做无依据低价承诺）", "区间", "10-21"),
    (14, "付款节点比例、质保期、备件要求", "商务条款", "惯例说明", "10-22"),
    (15, "立项评审/合同/进场/量产 期望日期", "项目里程碑", "日期", "10-21"),
    ("四、评估口径", None, None, None, None),
    (16, "是否已有其他供应商/方案在评估、评估维度与权重", "竞品分析与差异化", "文字", "10-21"),
    (17, "对标设备报价区间与产能（可披露部分）", "竞品定位（我方不外传）", "区间", "10-22"),
]

HEAD = ["序号", "资料项", "为什么需要（用途）", "期望格式", "期望时间", "贵方负责人", "贵方回复"]

wb = Workbook()
ws = wb.active
ws.title = "立项资料清单"

thin = Side(style="thin", color="C9D1D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="1F6FEB")
sec_fill = PatternFill("solid", fgColor="EAF1FB")

ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEAD))
t = ws.cell(row=1, column=1,
            value="荣旗 × Z-MAX 具身智能机器人项目 · 立项资料清单（2026-10-19）")
t.font = Font(name="微软雅黑", size=14, bold=True, color="0D1117")
t.alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[1].height = 26

ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEAD))
s = ws.cell(row=2, column=1,
            value="请于 10-21（周三）前回复已知项；未定项写「待现场确认」即可，我们按草案先行。"
                  "单台金额由我方按配置测算后报价，此处只需第 13 项目标预算/回收期。")
s.font = Font(name="微软雅黑", size=10, color="57606A")
s.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws.row_dimensions[2].height = 30

for i, h in enumerate(HEAD, 1):
    c = ws.cell(row=3, column=i, value=h)
    c.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    c.fill = head_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border
ws.row_dimensions[3].height = 24

r = 4
for row in ROWS:
    if row[1] is None:                                    # 分组标题行
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HEAD))
        c = ws.cell(row=r, column=1, value=row[0])
        c.font = Font(name="微软雅黑", size=10, bold=True, color="1F6FEB")
        c.fill = sec_fill
        c.alignment = Alignment(horizontal="left", vertical="center")
        for j in range(1, len(HEAD) + 1):
            ws.cell(row=r, column=j).border = border
        ws.row_dimensions[r].height = 22
        r += 1
        continue
    for j, v in enumerate(row + (None, None), 1):
        c = ws.cell(row=r, column=j, value=v)
        c.font = Font(name="微软雅黑", size=10)
        c.alignment = Alignment(horizontal="center" if j in (1, 5) else "left",
                               vertical="center", wrap_text=True)
        c.border = border
    ws.row_dimensions[r].height = 30
    r += 1

for j, w in enumerate([6, 46, 34, 14, 10, 12, 30], 1):
    ws.column_dimensions[get_column_letter(j)].width = w

ws.freeze_panes = "A4"

# ── 第二张表: 我方口径 (避免客户/商务来回问) ──
ws2 = wb.create_sheet("我方口径与时间")
NOTES = [
    ("项目", "荣旗项目（Z100Fd 立项）"),
    ("触发", "秦博 2026-10-19 周会 / 战略会：尽快发起立项"),
    ("资料截止", "2026-10-21（周三）"),
    ("技术方案 + 立项技术附件", "2026-10-21 提交（场景映射/四层架构/性能口径/交付清单/风险）"),
    ("完整立项包", "2026-10-23 提交（含单台配置与报价、销量预测、竞品分析）"),
    ("不需要客户提供", "单台金额（我方按 BOM 测算后报价）；只请客户给目标预算/回收期"),
    ("已有可引用", "76 功能点与能力质数签名 · KPI 口径（插入成功率≥99% · 节拍<12秒）· 12 件专利 · 7 任务全流程 · AOI 两条检测链"),
    ("我方不外传", "第 17 项竞品报价仅用于方案定位"),
]
ws2.cell(row=1, column=1, value="口径项").font = Font(name="微软雅黑", size=11, bold=True)
ws2.cell(row=1, column=2, value="内容").font = Font(name="微软雅黑", size=11, bold=True)
for i, (k, v) in enumerate(NOTES, 2):
    ws2.cell(row=i, column=1, value=k).font = Font(name="微软雅黑", size=10, bold=True)
    c = ws2.cell(row=i, column=2, value=v)
    c.font = Font(name="微软雅黑", size=10)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws2.row_dimensions[i].height = 28
ws2.column_dimensions["A"].width = 26
ws2.column_dimensions["B"].width = 96

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("✅ 已生成", OUT)
