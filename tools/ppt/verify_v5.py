# -*- coding: utf-8 -*-
"""verify_v5.py — 交付前一致性闸"""
import os
import re
import sys

from pptx import Presentation

BASE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(BASE, "Z-MAX_竞品与合作分析_v5_四表全量_2026-10-24.pptx")
prs = Presentation(DECK)
n = len(prs.slides._sldIdLst)
fail = []

print("页数: %d" % n)
if n != 33:
    fail.append("页数 %d != 33" % n)

ntab = npic = ntxt = 0
bad_star = []
bad_char = []
small = []
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_table:
            ntab += 1
        if sh.shape_type == 13:
            npic += 1
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if not t.strip():
            continue
        ntxt += 1
        if "**" in t:
            bad_star.append((i, t[:40]))
        for ch in ("✗", "✘", "✔", "□"):
            if ch in t:
                bad_char.append((i, ch, t[:30]))
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size and r.font.size.pt < 6.0:
                    small.append((i, r.font.size.pt, r.text[:24]))

print("表格: %d 个   图片: %d 张   文本块: %d 个" % (ntab, npic, ntxt))
print("`**` 残留: %d" % len(bad_star))
for x in bad_star:
    print("   P%d %s" % x)
    fail.append("P%d 有 ** 字面量" % x[0])
print("豆腐块字符: %d" % len(bad_char))
for x in bad_char:
    print("   P%d 含 %r  %s" % x)
    fail.append("P%d 含不可渲染字符" % x[0])
print("小于 6pt 的文字: %d 处" % len(small))
for x in small[:6]:
    print("   P%d %.1fpt %s" % x)

# 关键数字交叉核对（数据源 vs 文稿）
import json
D = json.load(open(os.path.join(BASE, "data_v5.json"), encoding="utf-8"))
hy = sum(1 for r in D["hybrid15"] if r[6])
print("源数据: 混合表有痕迹 %d / 15，模型表 %d 家" % (hy, len(D["model20"])))
alltext = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes
                   if sh.has_text_frame)
if "只有 9 家有数据痕迹" in alltext:
    fail.append("P25 文案仍是 9 家，与源数据 %d 家矛盾" % hy)

print()
if fail:
    print("FAIL %d 项:" % len(fail))
    for f in fail:
        print("  -", f)
    sys.exit(1)
print("PASS 全部通过")
