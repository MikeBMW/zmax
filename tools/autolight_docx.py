#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
autolight_docx.py — /autoLight 工艺卡 → Word(.docx) 一键导出

真源是页面里的 CARDS 数据（tools/web/autoLight.html），本模块只负责**排版**：
  前端 POST 上来 {cards:[{n,title,stage,act,det[],kv[[k,v]],tags[],fn,img}], meta:{...}}
  这里用 python-docx 生成封面 + 工艺路线总览 + 每步一张工艺卡（含三维视图截图）。

为什么要前端传内容：工艺卡文字只允许有一个真源（页面），服务端不复制一份，
避免"改了页面忘了改导出"。图片是 canvas.toDataURL('image/png') 的 data URL。
"""
from __future__ import annotations

import base64
import io
import re
import html as _html
import time

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CJK = "微软雅黑"
MONO = "Consolas"
TEAL = RGBColor(0x0F, 0x6E, 0x64)
GRAY = RGBColor(0x66, 0x66, 0x66)


# ---------------------------------------------------------------- 小工具
def _strip(s: str) -> str:
    """HTML 片段 → 纯文本（工艺卡正文里有 <b>/<span>）"""
    if s is None:
        return ""
    s = str(s)
    s = re.sub(r"<br\s*/?>", " ", s, flags=re.I)
    s = re.sub(r"<[^>]+>", "", s)
    return _html.unescape(s).replace("\u00a0", " ").strip()


def _font(run, size=None, bold=None, color=None, name=CJK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _para(doc, text, size=10.5, bold=False, color=None, space_after=4,
          align=None, style=None, name=CJK):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    _font(r, size, bold, color, name)
    return p


def _heading(doc, text, level=1):
    h = doc.add_heading("", level=level)
    r = h.add_run(text)
    _font(r, {1: 15, 2: 11.5}.get(level, 11), True, TEAL if level == 1 else None)
    h.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    h.paragraph_format.space_after = Pt(4)
    return h


def _bullet(doc, text, size=10):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    _font(r, size)
    return p


def _img_from_dataurl(s: str):
    if not s or not isinstance(s, str):
        return None
    m = re.match(r"^data:image/(?:png|jpeg|jpg);base64,(.+)$", s.strip(), re.S)
    if not m:
        return None
    try:
        return io.BytesIO(base64.b64decode(m.group(1)))
    except Exception:
        return None


def _shade(cell, hexcolor="F2F5F8"):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = tcPr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto",
                                        qn("w:fill"): hexcolor})
    tcPr.append(sh)


def _kv_table(doc, kv):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for k, v in kv:
        row = t.add_row().cells
        row[0].width = Cm(3.6)
        row[1].width = Cm(12.4)
        p0 = row[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(1)
        _font(p0.add_run(_strip(k)), 9.5, True)
        _shade(row[0])
        p1 = row[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(1)
        _font(p1.add_run(_strip(v)), 9.5)
    return t


def _overview_table(doc, cards):
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    hdr = ["序号", "工艺步骤", "阶段"]
    for i, h in enumerate(hdr):
        c = t.rows[0].cells[i]
        c.width = Cm(1.6 if i == 0 else (9.4 if i == 1 else 5.0))
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        _font(p.add_run(h), 9.5, True)
        _shade(c, "DCE8F0")
    for c in cards:
        row = t.add_row().cells
        vals = [c.get("n", ""), _strip(c.get("title", "")), _strip(c.get("stage", ""))]
        for i, v in enumerate(vals):
            p = row[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            _font(p.add_run(v), 9.5, i == 1)
    return t


# ---------------------------------------------------------------- 主函数
def build_docx(payload: dict) -> bytes:
    cards = (payload or {}).get("cards") or []
    if not cards:
        raise ValueError("payload.cards 为空：页面没有传工艺卡内容")

    meta = (payload or {}).get("meta") or {}
    doc = Document()

    # 页面设置 A4 + 常规页边距
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.9)

    # 正文默认字体（中文必须显式给 eastAsia，否则 Word 里回落成宋体/等线）
    st = doc.styles["Normal"]
    st.font.name = CJK
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)

    # ---------- 封面区 ----------
    _para(doc, "光纤端面检测及清理自动化设备", 19, True, TEAL, 2, WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, "工艺卡（技术协议 rev1.0 §2.2 九步工艺）", 13, True, None, 10,
          WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc,
          "甲方：苏州旭创科技有限公司　　乙方：苏州智蜂创元科技有限公司", 9.5, False, GRAY, 2,
          WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc,
          "坐标系与 /ar 三维视图一致：X 右 / Y 近 / Z 上，单位 m，地面 z=0　|　"
          "本体 3000×1500×1900 mm　|　操作 2 人　|　载具托盘缓存 ≤15 PCS　|　NG 仓 4 PCS",
          9, False, GRAY, 6, WD_ALIGN_PARAGRAPH.CENTER)

    _para(doc, "本文档由 /autoLight 工艺卡页面一键导出，内容与页面同源（同一份工艺卡数据）。"
               "每步一张卡：三维视图 + 动作分解 + 技术细节 + 参数口径 + 协议来源。"
               "三维视图为按协议模块关系的 1:1 示意，非施工图；协议未定项标「设计评审确认」。",
          9.5, False, None, 8, WD_ALIGN_PARAGRAPH.JUSTIFY)

    # ---------- 工艺路线总览 ----------
    _heading(doc, "工艺路线总览", 1)
    _overview_table(doc, cards)

    # ---------- 节拍口径 ----------
    _heading(doc, "节拍与性能口径（协议 2.1）", 1)
    _kv_table(doc, [
        ("8-FC", "非全端面自动对焦 2S + 测试 2S 条件下：CT ≤ 70s，UPH ≥ 51，OEE ≥ 85%（以所有接头测试一次通过为前提）"),
        ("16-FC", "全端面自动对焦 2S + 测试 2S 条件下：CT ≤ 134s，UPH ≥ 26，OEE ≥ 85%（以所有接头测试一次通过为前提）"),
        ("在线仪表", "自动端检仪 ASTCHECK PRO 200KC ×4；回损测试仪 RLM1412A-1FA-24 ×1；端面清洗机 Offsoon Pro ×3"),
        ("对接信号", "OK / NG / 完成（测试板卡→设备）；启动（设备→测试板卡）。板卡供电、显示器接口 → 设计评审确认"),
    ])

    # ---------- 每步工艺卡 ----------
    for c in cards:
        doc.add_page_break()
        no = _strip(c.get("n", ""))
        title = _strip(c.get("title", ""))
        stage = _strip(c.get("stage", ""))
        _heading(doc, "%s　%s%s" % (no, title, ("（" + stage + "）") if stage else ""), 1)

        img = _img_from_dataurl(c.get("img") or "")
        if img is not None:
            try:
                doc.add_picture(img, width=Cm(16.6))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                _para(doc, "图 %s-1　%s（三维窗口截图，坐标系同 /ar）" % (no, title),
                      8.5, False, GRAY, 8, WD_ALIGN_PARAGRAPH.CENTER)
            except Exception as e:  # 图片坏了不许整个导出失败
                _para(doc, "（视图截图嵌入失败：%s: %s）" % (type(e).__name__, e), 8.5, False, GRAY, 4)

        act = _strip(c.get("act", ""))
        if act:
            _heading(doc, "动作分解", 2)
            _para(doc, act, 10.5, False, None, 6, WD_ALIGN_PARAGRAPH.JUSTIFY)

        det = [d for d in (_strip(x) for x in (c.get("det") or [])) if d]
        if det:
            _heading(doc, "技术细节", 2)
            for d in det:
                _bullet(doc, d)

        kv = [(_strip(k), _strip(v)) for k, v in (c.get("kv") or []) if _strip(k)]
        if kv:
            _heading(doc, "参数与口径", 2)
            _kv_table(doc, kv)

        tags = [_strip(t) for t in (c.get("tags") or []) if _strip(t)]
        if tags:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            _font(p.add_run("关键词："), 9, True, GRAY)
            _font(p.add_run("　".join(tags)), 9, False, GRAY)

        fn = _strip(c.get("fn", ""))
        if fn:
            _para(doc, fn, 8.5, False, GRAY, 2)

    # ---------- 尾页信息 ----------
    doc.add_page_break()
    _heading(doc, "导出说明", 1)
    _kv_table(doc, [
        ("数据来源", "/autoLight 工艺卡页面（tools/web/autoLight.html）与本文档同源，改页面即改导出"),
        ("协议依据", "《光纤端面检测及清理设备技术协议》rev1.0：§2.1 技术指标 / §2.2 工艺流程九步 / §2.3 交互信号 / §3 硬件配置"),
        ("坐标系", "与 /ar 同一 base 世界系：X 右 / Y 近 / Z 上，单位 m，地面 z=0"),
        ("三维说明", "三维窗口为按协议模块关系的 1:1 示意，用于说明设备位置、物料流向与动作顺序；施工以 LAYOUT 图为准"),
        ("生成时间", meta.get("ts") or time.strftime("%Y-%m-%d %H:%M:%S")),
        ("导出入口", meta.get("url") or "/autoLight"),
    ])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_demo() -> bytes:
    """离线自检用：一张卡的最小样本（无图片）"""
    return build_docx({
        "meta": {"ts": "self-test", "url": "/autoLight"},
        "cards": [{
            "n": "1", "title": "待测物料缓存仓一", "stage": "上料 · 缓存",
            "act": "人工把载具放到<b>人工放料位</b> → <b>推料模组</b>推入缓存仓一",
            "det": ["载具承载 <b>1 PCS</b> MPO-8FC / MPO-16FC 光纤", "浮动机构 XY 单边 <b>0.5 mm</b>"],
            "kv": [["节拍", "人工上料不计入 CT"], ["缓存能力", "≤15 PCS"]],
            "tags": ["推料模组", "扫码绑定"],
            "fn": "来源：技术协议 rev1.0 §2.2 第 1 步",
        }],
    })


if __name__ == "__main__":
    import sys
    out = sys.argv[1] if len(sys.argv) > 1 else "/tmp/autolight_demo.docx"
    data = build_demo()
    with open(out, "wb") as f:
        f.write(data)
    print("wrote %s  %d bytes" % (out, len(data)))
