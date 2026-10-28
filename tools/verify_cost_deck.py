#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""判据：《双臂协作工位 · 三档目标成本分析》PPT 是不是真按老倪 2026-10-28 的口径出的。

老倪: 「成分分析；基础 硬件 15万，软件 7万；精细 硬件 20万，软件10万；专家 硬件 22万 软件 15万。
       按照这个给出成本分析；不要写销售价。重写 PPT」

查七段 (全部**真读 pptx**，不看代码注释)：
  ① 文件存在且在 zmax_data/deliverables
  ② 页数 = 10，每页都有标题文字
  ③ 三档目标成本逐档对账：硬件 15/20/22 万、软件 7/10/15 万、单套 22/30/37 万（从 cost_targets.json 现算后比对页面文本）
  ④ 不写销售价：全文禁止出现 售价 / 报价 / 利润空间 / 毛利 / 分成 / 回本 / 销售（一个都不许有）
  ⑤ 图表 ≥ 4（柱状/折线）
  ⑥ 现估 vs 目标 的降本缺口 = bom_capex − hw_cost 逐档出现在文字里
  ⑦ 每件成本 = (折旧+实施摊销+运营)/年产量 逐档出现在文字里（小数点后 4 位）

跑法: cd /home/ubuntu/zmax && gui-venv311/bin/python tools/verify_cost_deck.py
"""
import json
import os
import sys

ROOT = "/home/ubuntu/zmax"
sys.path.insert(0, os.path.join(ROOT, "tools"))
from cost_model import build_tiers                                             # noqa: E402
from targets_model import build_targets                                        # noqa: E402

PPT = "/home/ubuntu/zmax_data/deliverables/双臂协作工位_成本分析_三档目标成本.pptx"
BAN = ("售价", "报价", "利润空间", "毛利", "分成", "回本", "销售")
CN = {"A": "基础", "B": "精细", "C": "专家"}

PASS, FAIL = [], []


def chk(name, ok, detail=None):
    (PASS if ok else FAIL).append(name)
    print("[%s] %s%s" % ("PASS" if ok else "FAIL", name,
                        (" — " + str(detail)) if detail not in (None, "") else ""))


def wan(x):
    return float(x) / 1e4


def main():
    from pptx import Presentation
    d = build_tiers()
    T = {t["id"]: t for t in d["tiers"]}
    TG = {t["id"]: t for t in build_targets()["tiers"]}
    ids = ["A", "B", "C"]

    # ①
    chk("① PPT 在 %s" % PPT, os.path.exists(PPT), PPT if os.path.exists(PPT) else "缺失")
    if not os.path.exists(PPT):
        return 1
    prs = Presentation(PPT)

    # ②
    n = len(prs.slides._sldIdLst)
    chk("② 页数 = 10", n == 10, n)
    titles = []
    for i, sl in enumerate(prs.slides):
        txts = [sh.text_frame.text.strip() for sh in sl.shapes
                if sh.has_text_frame and sh.text_frame.text.strip()]
        titles.append(txts[0] if txts else "")
        chk("② 第 %d 页有标题" % (i + 1), bool(titles[-1]), (titles[-1] or "")[:24])

    # 全文（含表格单元格 + 图表类别/系列名）
    allt, charts = [], 0
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.has_text_frame:
                allt.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        allt.append(c.text)
            if getattr(sh, "has_chart", False) and sh.has_chart:
                charts += 1
                try:
                    pl = sh.chart.plots[0]
                    allt.extend([str(c) for c in pl.categories])
                    for ser in pl.series:
                        allt.append(str(ser.name))
                        for v in ser.values:
                            allt.append("%.4f" % float(v))
                except Exception:                                               # noqa: BLE001
                    pass
    txt = " ".join(allt)

    # ③
    for k in ids:
        t = TG[k]
        chk("③ %s版 硬件 %.0f 万 / 软件 %.0f 万 / 单套 %.0f 万 三处都在" %
            (CN[k], wan(t["hw_cost"]), wan(t["sw_cost"]), wan(t["cost_total"])),
            ("%.0f" % wan(t["hw_cost"])) in txt and ("%.0f" % wan(t["sw_cost"])) in txt
            and ("%.0f" % wan(t["cost_total"])) in txt,
            "hw %.0f sw %.0f total %.0f" % (wan(t["hw_cost"]), wan(t["sw_cost"]), wan(t["cost_total"])))
    chk("③ 三档单套成本 = 22 / 30 / 37 万 (老倪口述口径)",
        [round(wan(TG[k]["cost_total"])) for k in ids] == [22, 30, 37],
        [round(wan(TG[k]["cost_total"])) for k in ids])

    # ④
    hits = [b for b in BAN if b in txt]
    chk("④ 不写销售价：全文无 %s 字样" % "/".join(BAN), not hits, hits)

    # ⑤
    chk("⑤ 图表 ≥ 4", charts >= 4, charts)

    # ⑥
    for k in ids:
        v = "%.1f" % wan(TG[k]["gap_to_bom"])
        chk("⑥ %s版 降本缺口 %.1f 万 出现" % (CN[k], wan(TG[k]["gap_to_bom"])), v in txt, v)

    # ⑦
    for k in ids:
        v = "%.4f" % T[k]["unit_cost"]
        chk("⑦ %s版 每件成本 %s 元 出现" % (CN[k], v), v in txt, v)

    print("\n=== %d PASS / %d FAIL ===" % (len(PASS), len(FAIL)))
    if FAIL:
        print("失败项:")
        for f in FAIL:
            print("  -", f)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
