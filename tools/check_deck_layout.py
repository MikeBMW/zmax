#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""版式自检: 越界 / 表格出页 / 文本溢出估算 (CJK 按 1em/字宽)"""
import sys, unicodedata
from pptx import Presentation
from pptx.util import Emu

p = sys.argv[1]
prs = Presentation(p)
W, H = prs.slide_width / 914400, prs.slide_height / 914400
bad = 0
for si, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        l, t = sh.left / 914400, sh.top / 914400
        w, h = sh.width / 914400, sh.height / 914400
        if l < -0.02 or t < -0.02 or l + w > W + 0.02 or t + h > H + 0.02:
            print(f"  ✗ s{si} 越界 {sh.shape_type} {sh.name} l{l:.2f} t{t:.2f} r{l+w:.2f} b{t+h:.2f}")
            bad += 1
        if not sh.has_text_frame:
            continue
        need = 0.0
        for para in sh.text_frame.paragraphs:
            text = "".join(r.text for r in para.runs)
            if not text:
                continue
            sz = max((r.font.size.pt if r.font.size else 12) for r in para.runs) if para.runs else 12
            # CJK 字宽 = 1em, 半角 ≈ 0.55em
            em = sum(1.0 if unicodedata.east_asian_width(c) in "WF" else 0.55 for c in text)
            per_line = max(1.0, (w - 0.06) * 72 / sz)
            lines = max(1, -(-em // per_line))
            ls = para.line_spacing or 1.15
            need += lines * sz * ls * 1.18 / 72
        if need > h * 1.06 + 0.04:
            print(f"  ! s{si} 文本可能溢出 {sh.name}: 需 {need:.2f}in / 有 {h:.2f}in | {sh.text_frame.text[:34]!r}")
            bad += 1
print(f"检查完毕: {len(prs.slides._sldIdLst)} 页 · {bad} 处待看 · 画布 {W:.2f}×{H:.2f}in")
