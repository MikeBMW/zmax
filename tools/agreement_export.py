#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""协议导出器 — 《机器人数据平台合作与模型授权协议》Word (2026-10-10)

与 tools/data_governance.py **同源**: 附件里的数据分类/权属/导出范围/模型转售规则
全部读 config/platform/zmax_data_governance.json —— 协议写的和工具强制的必须是同一份
真相, 否则条款是纸面的。

用法:
  python3 tools/agreement_export.py            # 生成 .docx (+ 结构化 JSON)
  python3 tools/config_center.py agreement     # 同上 (配置中心内)
"""
import json
import os
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sor_export import footer_pagenum, h, page_setup, para, sha, table  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GOV = os.path.join(ROOT, "config", "platform", "zmax_data_governance.json")
OUT_ROOT = os.path.join(ROOT, "outputs", "agreements")
DOC_NO = "AGR-DP-ROBOT-" + time.strftime("%Y%m%d")
VERSION = "V1.0"
TITLE = "机器人数据平台合作与模型授权协议"
CN_FONT = "微软雅黑"

SECTIONS = [
    ("1. 定义与合作结构", [
        "本协议由平台方、供应商两方签署, 生产用户作为最终使用方在附件 B 中确认。",
        "合作结构: 平台方提供机器人硬件与数据平台; 供应商在平台数据上训练模型, 可将成品模型交付/转售给生产用户;"
        "平台方按约定参与收益分配。",
        "本协议采用组合签约方式: 《机器人硬件销售合同》(硬件交付) + 《数据平台使用与模型授权协议》(平台/数据/模型)。"
        "若只签一份主协议, 即本文件。",
        "以下所有附件与正文同等效力; 附件中的数据分类、权属、导出范围以平台治理真源为唯一准据 (附件 C)。",
    ]),
    ("2. 平台使用范围", [
        "场景范围: %(scenarios)s。超出清单的场景须另行书面授权。",
        "账号: %(accounts)s; 平台按账号计费与审计。",
        "期限: %(term)s。",
        "地域: %(region)s。",
        "并发训练任务上限: %(jobs)s 个。",
        "供应商不得将平台账号、令牌、SDK 提供给第三方使用或转借。",
    ]),
    ("3. 数据权属", [
        "平台采集/上传/存储的数据按 D0–D4 分类管理, 每类的权属、可导出性、脱敏要求、保留期见附件 A。",
        "D0 现场产线原始数据: 平台方与生产用户共有, 平台方受托处理; 不得导出到供应商环境 (技术上物理拒发)。",
        "D1 训练数据集: 平台方所有; 可在授权场景/账号/期限内下发给供应商用于训练。",
        "D2 模型权重与训练产物: 按第 4 章与附件 B 约定; 每一份权重必须登记训练数据指纹与权重哈希。",
        "D3 工程与客户信息 (含标定真值/密钥/客户资料): 平台方所有, 禁止导出; 现场所需参数由平台方按需生成最小集下发。",
        "D4 平台算法/SDK/训练框架: 平台方背景知识产权, 供应商仅获使用权。",
        "供应商不得将平台数据用于训练协议范围外的模型, 不得留存副本超过授权期限 (默认交付期满后 6 个月)。",
        "数据导出必须经平台同步工具 (附件 C), 工具会按类别自动拒发禁导出项并留审计; 绕过工具的导出视为违约。",
    ]),
    ("4. 模型权属", [
        "平台方基础权重 (VLA-T / Z-Flow / INTACT 及后续版本) 归平台方所有。",
        "供应商在平台数据上训练产生的适配层 (LoRA、动作头、检测头等) 归供应商所有。",
        "对上述适配层的转售、二次分发、对外授权, 按第 5 章执行; 每次转售须在模型台账登记。",
        "模型回传平台时须提供四项指纹: 训练数据类别清单、数据哈希清单、训练配置、权重 sha256; 缺项即拒收"
        " (平台侧校验不通过不得进入默认档/部署白名单)。",
    ]),
    ("5. 转售权限", [
        "是否允许转售给生产用户: %(resale_allowed)s。",
        "是否需平台方书面同意: %(resale_consent)s。",
        "是否允许二次授权/子许可给第三方: %(sub_license)s。",
        "转售前须完成: 模型台账登记 (模型 ID/权属/训练数据指纹/授权范围/有效期) + 平台方书面同意。",
        "转售模型只能部署于平台方设备白名单 (Orin/工控机等已授权设备)。",
    ]),
    ("6. 收益分配", [
        "模型销售分成: %(rev_model)s。",
        "平台使用费: %(rev_platform)s。",
        "数据服务费: %(rev_data)s。",
        "结算周期、发票与逾期处理由双方商务条款另行约定; 未约定前按平台台账记录的使用量为准。",
    ]),
    ("7. 知识产权与背景技术", [
        "平台方原有算法、SDK、数据集、训练框架、工程库与画布配置归平台方所有; 本协议不构成任何所有权转让。",
        "供应商自有模型、自有工艺知识、自有工具链归供应商所有。",
        "双方各自保证所提供内容不侵犯第三方权利; 因侵权主张导致的损失由提供方承担。",
        "供应商不得反编译、逆向工程或试图获取平台方背景技术的实现细节。",
    ]),
    ("8. 合规与安全", [
        "脱敏要求: %(desensitize)s。",
        "访问控制: %(access)s。",
        "日志留存: %(log_retention)s; 双方均有权查阅与本方相关的同步审计记录 (附件 D)。",
        "部署环境: %(deploy)s。",
        "安全事件: %(incident)s。",
        "供应商须配合平台方进行数据流向核查, 包括但不限于提供本地数据副本清单与销毁证明。",
    ]),
    ("9. 验收与 SLA", [
        "平台可用性: %(sla_avail)s。",
        "接口稳定性: %(sla_api)s。",
        "数据回传质量: %(sla_data)s。",
        "支持响应: %(sla_support)s。",
        "未达 SLA 的处理: 按双方约定的服务抵扣或整改期限执行; 连续两个月未达标, 任一方可要求重新谈判。",
    ]),
    ("10. 保密条款", [
        "保密信息包括: 模型参数与权重、训练数据及其统计特征、客户信息、生产现场数据、平台配置与工程真值、商务条款。",
        "保密期限: 自披露起 5 年; 涉及个人信息的按法定期限。",
        "例外: 已公开、合法获得、依法必须披露 (披露前通知对方)。",
        "违约方须赔偿并可被立即终止合作。",
    ]),
    ("11. 违约责任与终止", [
        "超范围使用平台 (账号共享/超场景/超地域/超期限): 立即停止 + 按平台使用费的 2 倍支付违约金。",
        "违规导出被禁类别数据或泄露数据: 立即终止 + 赔偿 + 自费销毁全部副本并提供证明。",
        "未登记即转售或二次授权: 立即终止授权 + 追偿已获收益。",
        "协议终止后: 供应商须在 15 日内销毁平台数据副本并出具证明; 已交付生产用户的模型按已生效的转售授权继续执行, 但不再获得平台数据更新。",
        "不可抗力、争议解决与适用法律按双方主合同执行。",
    ]),
    ("12. 附则", [
        "本协议自双方签字盖章之日起生效。",
        "附件 A/B/C/D 与本协议正文同等效力; 与治理真源冲突时以真源为准。",
        "任何影响数据权属、导出范围、模型转售、收益分配的变更, 须书面确认并同步更新治理真源后重新导出本协议。",
    ]),
]


def build():
    g = json.load(open(GOV, encoding="utf-8"))
    mt, ps, cp, sla = g["model_terms"], g["platform_scope"], g["compliance"], g["sla"]
    doc = page_setup(Document())
    footer_pagenum(doc, "%s · %s            " % (DOC_NO, VERSION))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    from sor_export import _font
    _font(p.add_run(TITLE), size=22, bold=True)
    para(doc, "文件编号: %s" % DOC_NO, size=11)
    para(doc, "版本号: %s" % VERSION, size=11)
    para(doc, "签约方: %s  |  %s" % (g["parties"]["平台方"], g["parties"]["供应商"]), size=11)
    para(doc, "最终使用方: %s" % g["parties"]["生产用户"], size=11)
    para(doc, "导出方式: 配置中心导出 (治理真源 %s, version %s)"
         % (os.path.relpath(GOV, ROOT), g["version"]), size=11)
    para(doc, "导出时间: %s" % time.strftime("%Y-%m-%d %H:%M:%S"), size=11)
    para(doc, "说明: 本协议的数据分类、权属与导出范围由平台治理真源驱动, 并由平台侧同步工具强制执行 (附件 C)。",
         size=10, italic=True)

    fmt = {
        "scenarios": "、".join(ps["scenarios"]), "accounts": ps["accounts"], "term": ps["term"],
        "region": ps["region"], "jobs": ps["concurrent_training_jobs"],
        "resale_allowed": "允许" if mt["resale"]["allowed"] else "不允许",
        "resale_consent": "需要" if mt["resale"]["requires_written_consent"] else "不需要",
        "sub_license": "不允许" if not mt["resale"]["sub_license"] else "允许",
        "rev_model": mt["revenue_share"]["model_sale"],
        "rev_platform": mt["revenue_share"]["platform_fee"],
        "rev_data": mt["revenue_share"]["data_service_fee"],
        "desensitize": "、".join(cp["desensitize"]), "access": cp["access_control"],
        "log_retention": cp["log_retention"], "deploy": cp["deploy_env"], "incident": cp["incident"],
        "sla_avail": sla["platform_availability"], "sla_api": sla["api_stability"],
        "sla_data": sla["data_return_quality"], "sla_support": sla["support_response"],
    }
    for title, items in SECTIONS:
        h(doc, title, 1)
        for it in items:
            para(doc, it % fmt if "%(" in it else it)

    doc.add_page_break()
    h(doc, "附件 A. 数据分类与权属表", 1)
    para(doc, "本表由治理真源生成; 平台同步工具按本表逐条强制 (可导出=false 的类别物理拒发)。", size=10)
    table(doc, ["类别", "名称", "权属", "可下发供应商", "脱敏要求", "保留期"],
          [[c["class_id"], c["name"], c["owner"], "是" if c["exportable_to_supplier"] else "否 (拒发)",
            "、".join(c["desensitize"]) or "无", c["retention"]] for c in g["data_classes"]],
          widths=[1.2, 3.2, 3.4, 2.0, 4.4, 2.0])
    para(doc, "", size=6)
    for c in g["data_classes"]:
        para(doc, "· %s %s — %s" % (c["class_id"], c["name"], c["note"]), size=9)

    h(doc, "附件 B. 模型权属与转售规则", 1)
    table(doc, ["项", "约定"], [
        ["平台方基础权重", mt["base_weights"]],
        ["供应商适配层", mt["finetuned_layers"]],
        ["共有/衍生", mt["joint_ownership"]],
        ["转售", "允许=%s · 需书面同意=%s · 二次授权=%s" % (mt["resale"]["allowed"],
                                                            mt["resale"]["requires_written_consent"],
                                                            mt["resale"]["sub_license"])],
        ["回传必附指纹", "、".join(mt["fingerprint_required"])],
    ], widths=[3.6, 12.4])
    para(doc, "生产用户确认栏: 授权范围 ______  有效期 ______  设备白名单 ______  签章 ______", size=10)

    h(doc, "附件 C. 数据同步机制 (技术附件)", 1)
    para(doc, "本附件与平台工具同源: 协议条款由 tools/data_governance.py 强制执行, 不可绕过。", size=10)
    table(doc, ["通道", "名称", "方向", "传输与校验", "闸门"],
          [[c["id"], c["name"], c["direction"], "%s | %s" % (c["transport"], c["integrity"]), c["gate"]]
           for c in g["sync_channels"]], widths=[1.2, 2.2, 3.0, 5.6, 4.0])
    para(doc, "可执行命令 (平台侧):", size=10, bold=True)
    for c in [
        "python3 tools/data_governance.py plan supplier          # 下发计划: 允许/拦截清单",
        "python3 tools/data_governance.py manifest --class D1    # 生成带权属标签与哈希的清单",
        "python3 tools/data_governance.py push supplier --class D1 --to <目录>   # 真下发, 被禁类别物理拒发",
        "python3 tools/data_governance.py pull supplier --from <回传目录>          # 登记台账, 缺指纹/哈希不符拒收",
        "python3 tools/data_governance.py ledger | audit -n 20   # 模型台账 / 同步审计",
    ]:
        para(doc, "   " + c, size=9)

    h(doc, "附件 D. 审计与留存", 1)
    table(doc, ["项", "值"], [
        ["审计日志", "data/governance/audit.jsonl (逐条: 时间/主体/动作/对象/结果)"],
        ["模型台账", "data/governance/model_ledger.json (模型/权属/训练数据指纹/权重哈希/授权范围/转售状态)"],
        ["数据清单", "data/governance/manifests/ (带权属标签与 sha256)"],
        ["留存期", cp["log_retention"]],
        ["治理真源", "%s  version %s  sha256=%s…" % (os.path.relpath(GOV, ROOT), g["version"], sha(GOV)[:16])],
        ["导出时间", time.strftime("%Y-%m-%d %H:%M:%S")],
    ], widths=[4.0, 12.0])
    return doc, g


def main():
    doc, g = build()
    ts = time.strftime("%Y%m%d_%H%M%S")
    out = os.path.join(OUT_ROOT, ts)
    os.makedirs(out, exist_ok=True)
    base = "%s_%s" % (DOC_NO, VERSION)
    fp = os.path.join(out, base + ".docx")
    doc.save(fp)
    jp = os.path.join(out, base + ".json")
    json.dump({"doc_no": DOC_NO, "version": VERSION, "title": TITLE,
               "governance": g, "sections": [{"title": t, "clauses": i} for t, i in SECTIONS]},
              open(jp, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    man = {"doc_no": DOC_NO, "version": VERSION, "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
           "source_governance": os.path.relpath(GOV, ROOT), "governance_sha256": sha(GOV),
           "source_db": "data/database/zmax/zmax_engineering.db",
           "source_db_sha256": sha(os.path.join(ROOT, "data", "database", "zmax", "zmax_engineering.db")),
           "files": {os.path.basename(fp): {"bytes": os.path.getsize(fp), "sha256": sha(fp)},
                     os.path.basename(jp): {"bytes": os.path.getsize(jp), "sha256": sha(jp)}}}
    json.dump(man, open(os.path.join(out, "manifest.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    print("✅ 协议 (Word): %s" % fp)
    print("   大小 %.1f KB · 正文 12 章 + 附件 A(数据权属)/B(模型与转售)/C(同步机制)/D(审计)"
          % (os.path.getsize(fp) / 1024.0))
    print("   目录: %s" % out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
